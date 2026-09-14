"""
Opakující se bloky: výpis stringů a tabulka stringů.

Tohle se pouhým přepisem hodnot vyřešit nedá. Šablona MŠ Tichá má čtyři
stringy po 25 panelech; nový projekt může mít tři po 14. Musí se tedy
změnit i **počet** odstavců a řádků tabulky, ne jen čísla v nich.

Řádek se nikdy nevyrábí od nuly – naklonuje se existující řádek šablony
a přepíše se jeho obsah. Tím zůstane zachované formátování, ohraničení
i výška řádku.
"""
from __future__ import annotations

import copy
import re
from typing import Any

from docx.table import Table
from docx.text.paragraph import Paragraph

# Výpis stringů ve V1: „25 ks FV panelů– string 1.1 <tab>– střídač 33,3 kW
# (INV1) – 25 ks optimizeru S500“
STRING_LINE = re.compile(
    r"^\s*(?P<n>\d+)\s*ks\s*FV\s*panelů[–\-]\s*string\s*(?P<id>[\d.]+)", re.I
)
STRING_TABLE_HEADER = re.compile(r"string\s*č\.?", re.I)


def cz(value: Any) -> str:
    """Číslo v českém zápisu – desetinná čárka, bez zbytečné nuly."""
    if isinstance(value, bool) or value is None or value == "":
        return ""
    if isinstance(value, (int, float)):
        return f"{value:g}".replace(".", ",")
    return str(value)


# Krytiny ve tvaru „s krytinou z …“. Co tu není, zůstane v prvním pádě
# a generátor to ohlásí – lepší než vymyslet špatný tvar.
ROOF_GENITIVE = {
    "asfaltové pásy": "asfaltových pásů",
    "asfaltový pás": "asfaltového pásu",
    "trapézový plech": "trapézového plechu",
    "falcovaný plech": "falcovaného plechu",
    "plechová krytina": "plechové krytiny",
    "betonová taška": "betonové tašky",
    "pálená taška": "pálené tašky",
    "vlnitý eternit": "vlnitého eternitu",
    "pvc folie": "PVC folie",
    "fólie": "fólie",
}


def roof_genitive(name: str) -> tuple[str, bool]:
    """Vrací tvar po předložce „z“ a příznak, zda se povedlo skloňovat."""
    key = (name or "").strip().lower()
    if key in ROOF_GENITIVE:
        return ROOF_GENITIVE[key], True
    return name, False


def string_line_text(item: dict[str, Any], pattern: str) -> str:
    """Naplní vzor jedné řádky výpisu stringů daty jednoho stringu."""
    return (
        pattern.replace("{n}", str(item["panel_count"]))
        .replace("{id}", str(item["id"]))
        .replace("{inv_p}", cz(item.get("inverter_p_ac", "")))
        .replace("{inv}", str(item.get("inverter", "")))
        .replace("{opt}", str(item.get("optimizer", "")))
    )


def derive_line_pattern(sample: str) -> str:
    """
    Ze vzorové řádky šablony udělá vzor s místy k doplnění.
    Zachová interpunkci, tabulátor i mezery přesně tak, jak je psaná.
    """
    m = STRING_LINE.match(sample)
    if not m:
        return sample
    out = sample
    out = re.sub(r"\(INV\d+\)", "({inv})", out)
    out = re.sub(r"střídač\s*[\d.,]+\s*kW", "střídač {inv_p} kW", out, flags=re.I)
    out = re.sub(r"optimizeru\s*\S+", "optimizeru {opt}", out, flags=re.I)
    out = re.sub(r"string\s*[\d.]+", "string {id}", out, flags=re.I)
    out = re.sub(r"\b\d+\s*ks", "{n} ks", out)
    return out


def find_string_paragraphs(doc) -> list[tuple[str, Paragraph]]:
    """Odstavce výpisu stringů, i s cestou, aby je uměl najít generátor."""
    from .render import build_path_index

    found: list[tuple[str, Paragraph]] = []
    for path, par in build_path_index(doc).items():
        if "/tbl" in path:
            continue
        if STRING_LINE.match(par.text or ""):
            found.append((path, par))
    return found


def find_string_table(doc) -> Table | None:
    for table in doc.tables:
        if not table.rows:
            continue
        head = " ".join(c.text for c in table.rows[0].cells)
        if STRING_TABLE_HEADER.search(head):
            return table
    return None


def _set_paragraph_text(par: Paragraph, text: str) -> None:
    """Přepíše odstavec a nechá formátování prvního runu."""
    runs = [r for r in par.runs if r.text]
    if not runs:
        return
    runs[0].text = text
    for r in runs[1:]:
        r.text = ""


def _delete(element) -> None:
    element.getparent().remove(element)


def _clone_after(element):
    new = copy.deepcopy(element)
    element.addnext(new)
    return new


def apply_string_list(doc, strings: list[dict[str, Any]]) -> int:
    """
    Srovná počet odstavců výpisu stringů s počtem stringů v projektu
    a naplní je. Vrací počet zapsaných řádek.
    """
    found = find_string_paragraphs(doc)
    if not found or not strings:
        return 0
    pattern = derive_line_pattern(found[0][1].text)
    pars = [p for _, p in found]

    while len(pars) < len(strings):
        new_el = _clone_after(pars[-1]._p)
        pars.append(Paragraph(new_el, pars[-1]._parent))
    for extra in pars[len(strings) :]:
        _delete(extra._p)
    pars = pars[: len(strings)]

    for par, item in zip(pars, strings):
        _set_paragraph_text(par, string_line_text(item, pattern))
        for run in par.runs:
            from .docx_model import clear_highlight

            clear_highlight(run)
    return len(pars)


def apply_string_table(doc, strings: list[dict[str, Any]]) -> int:
    """
    Naplní tabulku stringů. Sloupce se berou podle záhlaví, ne podle
    pořadí – tabulka se mezi sadami liší.
    """
    table = find_string_table(doc)
    if table is None or not strings:
        return 0

    head = [c.text.strip().lower() for c in table.rows[0].cells]

    def col(*names: str) -> int | None:
        for i, h in enumerate(head):
            if any(n in h for n in names):
                return i
        return None

    c_id = col("string")
    c_cnt = col("počet")
    c_type = col("typ panelu", "typ")
    c_pwr = col("výkon")
    c_u = col("napětí")
    c_i = col("proud")

    body = list(table.rows)[1:]
    while len(body) < len(strings):
        new_el = _clone_after(body[-1]._tr)
        table._tbl.remove(new_el)
        body[-1]._tr.addnext(new_el)
        body.append(table.rows[len(body) + 1])
    for extra in body[len(strings) :]:
        _delete(extra._tr)
    body = list(table.rows)[1 : 1 + len(strings)]

    from .docx_model import clear_highlight

    for row, item in zip(body, strings):
        values = {
            c_id: str(item["id"]),
            c_cnt: f"{item['panel_count']} ks.",
            c_type: str(item.get("panel_type", "")),
            c_pwr: str(item.get("power", "")),
            c_u: str(item.get("voltage", "")),
            c_i: str(item.get("current", "")),
        }
        for ci, text in values.items():
            if ci is None or ci >= len(row.cells) or not text:
                continue
            cell = row.cells[ci]
            if not cell.paragraphs:
                continue
            _set_paragraph_text(cell.paragraphs[0], text)
            for p in cell.paragraphs:
                for run in p.runs:
                    clear_highlight(run)
    return len(body)


def compute_strings(data: dict[str, Any]) -> list[dict[str, Any]]:
    """
    Doplní ke stringům odvozené hodnoty, které tabulka potřebuje.
    Nic se nehádá – co nejde spočítat z dodaných dat, zůstane prázdné
    a generátor to ohlásí (bod 29).
    """
    pv = data.get("pv", {})
    panel = pv.get("panel", {})
    inverters = {i.get("label"): i for i in pv.get("inverters", [])}
    out: list[dict[str, Any]] = []
    for s in pv.get("strings", []):
        n = s["panel_count"]
        inv = inverters.get(s.get("inverter"), {})
        wp = panel.get("wp")
        vmp = panel.get("vmp")
        imp = panel.get("imp")
        out.append(
            {
                "id": s["id"],
                "panel_count": n,
                "inverter": s.get("inverter", ""),
                "inverter_p_ac": inv.get("p_ac_kw", ""),
                "optimizer": pv.get("optimizers", {}).get("model", ""),
                "panel_type": panel.get("manufacturer", ""),
                "power": f"{n * wp / 1000:.2f}".replace(".", ",") + " kWp" if wp else "",
                "voltage": f"{n * vmp:.1f}".replace(".", ",") + " V" if vmp else "",
                "current": (f"{imp:.2f}".replace(".", ",") + " A") if imp else "",
            }
        )
    return out


# --------------------------------------------------------------- orientace

ORIENT_LINE = re.compile(r"\d+\s*panelů\s+je\s+orientováno", re.I)

# 8 světových stran podle azimutu (Q-08)
_DIRECTIONS = [
    (0, "sever"),
    (45, "severovýchod"),
    (90, "východ"),
    (135, "jihovýchod"),
    (180, "jih"),
    (225, "jihozápad"),
    (270, "západ"),
    (315, "severozápad"),
]


def direction_for(azimuth: float) -> str:
    """Azimut → světová strana. Sever = 0°, po 45°."""
    idx = int((float(azimuth) % 360 + 22.5) // 45) % 8
    return _DIRECTIONS[idx][1]


def _fmt_deg(value: Any) -> str:
    s = f"{float(value):g}".replace(".", ",")
    return s


def orientation_sentence(data: dict[str, Any], strings: list[dict[str, Any]]) -> str | None:
    """
    Věta o orientaci panelů, složená z dat projektu.

    V šabloně je tenhle odstavec zdrojem chyby: u MŠ Tichá tvrdí
    „42 panelů … a 42 panelů …“, ačkoli projekt má 100 panelů. Když se
    věta skládá z pv.orientation, součet sedí vždy.
    """
    pv = data.get("pv", {})
    items = pv.get("orientation") or []
    if not items:
        return None
    parts = [
        "{} panelů je orientováno na {} s azimutem {}° (sever = 0°)".format(
            it["count"], it.get("direction") or direction_for(it["azimuth"]),
            _fmt_deg(it["azimuth"]),
        )
        for it in items
    ]
    sentence = " a ".join(parts) + "."
    building = data.get("site", {}).get("building", {})
    tilt = building.get("roof_tilt_deg")
    covering = building.get("roof_covering")
    if strings and tilt is not None and covering:
        ids = ", ".join(str(s["id"]) for s in strings)
        form, ok = roof_genitive(covering)
        sentence += (
            "\nPanely ve stringu {} budou umístěny na ploché střeše "
            "se sklonem {}° s krytinou z {}.".format(ids, _fmt_deg(tilt), form)
        )
        if not ok:
            sentence += "  "
    return sentence


def apply_orientation(doc, data: dict[str, Any], strings: list[dict[str, Any]]) -> int:
    from .docx_model import clear_highlight
    from .render import build_path_index

    text = orientation_sentence(data, strings)
    if not text:
        return 0
    written = 0
    for path, par in build_path_index(doc).items():
        if "/tbl" in path or not ORIENT_LINE.search(par.text or ""):
            continue
        _set_paragraph_text(par, text)
        for run in par.runs:
            clear_highlight(run)
        written += 1
    return written


# ------------------------------------------------- výčty označení stringů

# „stringu 1.1 až 1.4“, „stringů 1.1 a 1.2“, „ve stringu 1.1, 1.2, 1.3, 1.4“
STRING_ENUM = re.compile(
    r"(?P<first>\d+\.\d+)"
    r"(?P<rest>(?:\s*(?:až|a|,)\s*\d+\.\d+)+)"
)


def format_string_ids(ids: list[str], style: str) -> str:
    """Výčet označení stringů ve stejném stylu, jakým je psaný v šabloně."""
    if not ids:
        return ""
    if len(ids) == 1:
        return ids[0]
    if style == "range":
        return f"{ids[0]} až {ids[-1]}"
    if len(ids) == 2:
        return f"{ids[0]} a {ids[1]}"
    return ", ".join(ids)


def apply_string_enumerations(doc, strings: list[dict[str, Any]]) -> int:
    """
    Srovná výčty označení stringů v textu se skutečnou konfigurací.

    V šabloně MŠ Tichá je tenhle údaj rozporný sám se sebou: kapitola 2.6
    mluví o stringech 1.1 až 1.4, kapitola 3.3 o 1.1 až 1.2, přestože
    stringy jsou čtyři. Když se výčet skládá z dat, rozpor nevznikne.
    """
    from .docx_model import clear_highlight
    from .render import build_path_index

    ids = [str(s["id"]) for s in strings]
    if not ids:
        return 0
    changed = 0
    for path, par in build_path_index(doc).items():
        text = par.text or ""
        if "string" not in text.lower():
            continue
        m = STRING_ENUM.search(text)
        if not m:
            continue
        style = "range" if "až" in m.group("rest") else "list"
        new_enum = format_string_ids(ids, style)
        if m.group(0) == new_enum:
            continue
        _set_paragraph_text(par, text[: m.start()] + new_enum + text[m.end() :])
        for run in par.runs:
            clear_highlight(run)
        changed += 1
    return changed


# ------------------------------------------------- DC řádky tabulky kabeláže

DC_CABLE_ROW = re.compile(r"^\s*WL\s*(?P<id>\d+\.\d+)\s*\((?P<pole>[+\-−])\)", re.I)


def apply_dc_cable_rows(doc, strings: list[dict[str, Any]]) -> int:
    """
    Srovná DC řádky tabulky kabeláže se stringy.

    Každý string má dva kabely, kladný a záporný. Když se změní počet
    stringů, musí se změnit i počet řádků – jinak by v tabulce zůstal
    kabel ke stringu, který v projektu není.
    """
    from .docx_model import clear_highlight

    if not strings:
        return 0
    written = 0
    for table in doc.tables:
        rows = list(table.rows)
        idx = [i for i, r in enumerate(rows) if r.cells and DC_CABLE_ROW.match(r.cells[0].text)]
        if not idx:
            continue
        first, last = idx[0], idx[-1]
        wanted = len(strings) * 2
        have = last - first + 1

        while have < wanted:
            _clone_after(table.rows[first + have - 1]._tr)
            have += 1
        for i in range(wanted, have):
            _delete(table.rows[first + wanted]._tr)

        rows = list(table.rows)
        for n, item in enumerate(strings):
            for k, pole in enumerate(("+", "-")):
                row = rows[first + n * 2 + k]
                sid = str(item["id"])
                _set_cell(row, 0, f"WL{sid}({pole}) ")
                for ci, cell in enumerate(row.cells):
                    if "string" in cell.text.strip().lower():
                        _set_cell(row, ci, f"String {sid} ")
                for cell in row.cells:
                    for p in cell.paragraphs:
                        for run in p.runs:
                            clear_highlight(run)
                written += 1
    return written


def _set_cell(row, index: int, text: str) -> None:
    if index >= len(row.cells):
        return
    cell = row.cells[index]
    if not cell.paragraphs:
        return
    _set_paragraph_text(cell.paragraphs[0], text)
