"""
Kontrola hotového dokumentu a odolnost celého enginu.

`test_all_templates_survive` projede všech 30 dodaných souborů – šablony
i hotové projekty. Je to pojistka proti tomu, aby jeden neobvyklý dokument
shodil generování.
"""
from __future__ import annotations

import glob
import os
from pathlib import Path

import docx
import pytest

from app.binder import rewrite_free_text, rewrite_unmarked
from app.parser import parse_template
from app.render import render, values_from_manifest
from app.sequences import compute_strings
from app.verify import verify

TEMPLATES = Path(os.environ.get("FVE_TEMPLATES", "data/templates"))
V1_TZ = TEMPLATES / "V1" / "D.2.2-TZ-FVE-MS Tichá.docx"
ALL_DOCS = sorted(glob.glob(str(TEMPLATES / "**" / "*.docx"), recursive=True))

pytestmark = pytest.mark.skipif(not V1_TZ.exists(), reason="šablony nejsou k dispozici")

PROJECT = {
    "pv.count": "40 ks",
    "pv.panel.wp": "550 Wp",
    "pv.panel.model": "AIKO Comet 3N",
    "energy.p_dc_kwp": "22,0 kWp",
    "pv.inverters.0.model": "SUN-20K-SG01HP3",
    "pv": {
        "panel": {"manufacturer": "AIKO", "wp": 550, "vmp": 41.2, "imp": 13.21},
        "inverters": [{"label": "INV1", "p_ac_kw": 20.0}],
        "optimizers": {"model": "S500B"},
        "orientation": [{"azimuth": 95, "count": 20}, {"azimuth": 275, "count": 20}],
        "strings": [
            {"id": "1.1", "panel_count": 14, "inverter": "INV1"},
            {"id": "1.2", "panel_count": 13, "inverter": "INV1"},
            {"id": "1.3", "panel_count": 13, "inverter": "INV1"},
        ],
    },
    "site": {"building": {"roof_tilt_deg": 3, "roof_covering": "Asfaltové pásy"}},
}


def _full_render(template: Path, out: Path, *, rewrite_fixed: bool = True):
    manifest = parse_template(template)
    values = values_from_manifest(manifest, PROJECT)
    free_text, _ = rewrite_free_text(manifest, PROJECT)
    values.update(free_text)
    if rewrite_fixed:
        unmarked, _ = rewrite_unmarked(manifest, PROJECT)
        values.update(unmarked)
    render(
        template,
        out,
        values,
        strings=compute_strings(PROJECT),
        project=PROJECT,
    )
    return manifest


def test_generated_document_has_no_leftovers(tmp_path: Path) -> None:
    out = tmp_path / "out.docx"
    manifest = _full_render(V1_TZ, out)
    report = verify(out, manifest, PROJECT)
    assert report.ok, [
        f"{lo.key}: {lo.value!r} v {lo.where} — {lo.context}" for lo in report.leftovers
    ]


def test_verify_catches_forgotten_value(tmp_path: Path) -> None:
    """
    Bez přepisu neoznačeného textu zůstane ve větě panel z Tiché.
    Kontrola to musí najít – jinak by byla k ničemu.
    """
    out = tmp_path / "out.docx"
    manifest = _full_render(V1_TZ, out, rewrite_fixed=False)
    report = verify(out, manifest, PROJECT)
    assert not report.ok
    assert any(lo.key == "pv.panel.model" for lo in report.leftovers)


def test_verify_ignores_norm_citation(tmp_path: Path) -> None:
    """
    „ochrana … do 1000 V na straně AC“ je citace normy, i když 1000 V
    je zároveň parametr střídače. Nesmí se hlásit jako zapomenutá hodnota.
    """
    out = tmp_path / "out.docx"
    manifest = _full_render(V1_TZ, out)
    report = verify(out, manifest, PROJECT)
    assert not any("1000 V" in lo.context for lo in report.leftovers)


@pytest.mark.parametrize("path", ALL_DOCS, ids=[Path(p).stem[:28] for p in ALL_DOCS])
def test_all_templates_survive(path: str, tmp_path: Path) -> None:
    """Import i generování projde na každém dodaném dokumentu."""
    out = tmp_path / "out.docx"
    _full_render(Path(path), out)
    assert out.exists() and out.stat().st_size > 10_000
    docx.Document(str(out))  # soubor musí jít znovu otevřít


def test_render_is_idempotent(tmp_path: Path) -> None:
    """
    Dvojí generování ze stejných dat dá stejný dokument.
    Kdyby klonování řádků přidávalo pokaždé další, projeví se to tady.
    """
    a, b = tmp_path / "a.docx", tmp_path / "b.docx"
    _full_render(V1_TZ, a)
    _full_render(V1_TZ, b)

    def shape(p: Path) -> tuple:
        d = docx.Document(str(p))
        return (
            len(d.paragraphs),
            len(d.tables),
            tuple(len(t.rows) for t in d.tables),
            "\n".join(x.text for x in d.paragraphs),
        )

    assert shape(a) == shape(b)


def test_regenerating_from_output_is_stable(tmp_path: Path) -> None:
    """
    Vygenerovaný dokument použitý jako vstup nesmí znovu měnit počty
    řádků. Tohle chytá případ, kdy si uživatel uloží výstup a ten se
    omylem dostane zpátky do generování.
    """
    first, second = tmp_path / "1.docx", tmp_path / "2.docx"
    _full_render(V1_TZ, first)
    _full_render(first, second)
    d1, d2 = docx.Document(str(first)), docx.Document(str(second))
    assert [len(t.rows) for t in d1.tables] == [len(t.rows) for t in d2.tables]
    assert len(d1.paragraphs) == len(d2.paragraphs)
